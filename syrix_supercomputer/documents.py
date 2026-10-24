"""Real files: documents, spreadsheets and PDFs in the task workspace.

"Create a spreadsheet from this video" has to produce a file that opens in
Excel, not a markdown table in the chat. Each format is produced by something
that actually writes that format:

======  ===========================================================
md/txt  written as-is
html    a real template + theme, and it is also the source for pdf
csv     stdlib :mod:`csv`, so quoting and embedded commas are right
xlsx    openpyxl — formats, frozen header, autofilter, optional chart
docx    python-docx — headings, tables, inline bold/italic, images
pptx    python-pptx
pdf     Chromium's own print engine, via the embedded browser
======  ===========================================================

PDF deliberately has no Python dependency. The HTML is written into the
workspace, served by the same local preview server the coding workspace uses,
opened in a Supercomputer tab and printed with ``webContents.printToPDF``. That
is the one PDF renderer already on this machine that is guaranteed to match what
The user saw on screen, and it needs nothing installed.

2026-07-26 — templates, themes and images
-----------------------------------------
Everything above the print engine was rebuilt. The verdict on the old output was
that it was *lazy*: no structure, no emphasis, nothing that looked designed. It
was, because there was exactly one hardcoded stylesheet and every document went
through it — a CV came out looking like a memo.

Now :mod:`syrix_supercomputer.doc_templates` supplies ten real templates and
:mod:`syrix_supercomputer.doc_themes` six colour schemes, chosen independently.
Images (a CV photo, a generated chart) are embedded as data URIs, because the
print runs against a local origin with no network and anything remote would
silently fail to render.

Every path goes through :mod:`syrix_supercomputer.workspace`, so containment,
the refused-suffix list and the size cap apply to documents exactly as they do
to code.
"""
from __future__ import annotations

import base64
import csv
import html as _html
import io
import os
import re
import time
from typing import Any, Dict, List, Optional

import syrix_obs as obs

from . import doc_templates, doc_themes, workspace

FORMATS = ("md", "txt", "html", "csv", "xlsx", "docx", "pptx", "pdf")

# Images small enough to embed. A CV photo is tens of KB; a 4 MB camera JPEG
# base64-encodes to 5.3 MB of HTML and blows the workspace file cap, so it is
# downscaled rather than refused.
MAX_IMAGE_BYTES = 4 * 1024 * 1024
IMAGE_MAX_EDGE = 900

# The single hardcoded shell that used to live here has been replaced by
# doc_templates.render(). It is kept only as the fallback for a raw HTML
# fragment with no template context, and is deliberately plain.
_FALLBACK_SHELL = """<!doctype html>
<html><head><meta charset="utf-8"><title>{title}</title><style>
  @page {{ margin: 18mm 16mm; }}
  body {{ font: 11.5pt/1.55 -apple-system, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
          color: #16181d; margin: 0; }}
  h1 {{ font-size: 21pt; margin: 0 0 10pt; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #d5d8dd; padding: 5pt 7pt; text-align: left; }}
</style></head><body>
<h1>{title}</h1>
{body}
</body></html>
"""


def _slug(name: str) -> str:
    """A safe, readable filename stem.

    Separators are dropped rather than escaped, so a model asking for
    ``../../escape`` gets a file in the workspace rather than a refusal — but
    the leading dots go too, because ``.....escape.md`` is a hidden-looking
    file nobody asked for. Containment itself is still enforced by
    :func:`workspace.resolve`; this is about the name being sane.
    """
    base = re.sub(r"[^A-Za-z0-9._ -]+", "", str(name or "").strip())
    base = re.sub(r"\.{2,}", ".", base).strip(". -_")
    base = re.sub(r"\s+", "_", base) or "document"
    return base[:80]


def filename_for(name: str, fmt: str) -> str:
    """A workspace-relative filename with exactly one correct extension."""
    stem = _slug(name)
    ext = f".{fmt}"
    if stem.lower().endswith(ext):
        return stem
    return re.sub(r"\.[A-Za-z0-9]{1,5}$", "", stem) + ext


# ── markdown → html ────────────────────────────────────────────────────────
def markdown_to_html(text: str) -> str:
    """Render markdown for the HTML/PDF path.

    Uses the `markdown` package when present (it is, in this venv) and falls
    back to a small converter rather than emitting raw markdown into a PDF —
    a missing optional dependency should degrade the typography, not the file.
    """
    body = str(text or "")
    try:
        import markdown as _md

        return _md.markdown(body, extensions=["tables", "fenced_code", "sane_lists"])
    except Exception as exc:
        obs.report("supercomputer.documents.markdown", exc, level="debug")

    out: List[str] = []
    for block in re.split(r"\n{2,}", body):
        chunk = block.strip()
        if not chunk:
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", chunk)
        if heading:
            lvl = len(heading.group(1))
            out.append(f"<h{lvl}>{_html.escape(heading.group(2))}</h{lvl}>")
            continue
        if all(re.match(r"^\s*[-*]\s+", ln) for ln in chunk.splitlines()):
            bullets = [_html.escape(re.sub(r"^\s*[-*]\s+", "", ln))
                       for ln in chunk.splitlines()]
            items = "".join(f"<li>{b}</li>" for b in bullets)
            out.append(f"<ul>{items}</ul>")
            continue
        para = _html.escape(chunk).replace("\n", "<br>")
        out.append(f"<p>{para}</p>")
    return "\n".join(out)


def wrap_html(title: str, body_html: str, *, template: str = "", theme: str = "",
              subtitle: str = "", photo: str = "",
              meta: Optional[Dict[str, Any]] = None,
              config: Optional[Dict[str, Any]] = None) -> str:
    """A complete document. Goes through the template system unless it cannot."""
    try:
        return doc_templates.render(
            body_html, template=template, theme=theme, title=title,
            subtitle=subtitle, photo=photo, meta=meta, config=config,
        )
    except Exception as exc:
        obs.report("supercomputer.documents.template", exc, template=template)
        return _FALLBACK_SHELL.format(
            title=_html.escape(str(title or "Document")), body=body_html)


# ── images ─────────────────────────────────────────────────────────────────
_IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp"}


def image_data_uri(path: str, *, max_edge: int = IMAGE_MAX_EDGE) -> str:
    """An image as a ``data:`` URI, downscaled if it is large. "" on failure.

    A data URI rather than a file reference because the print engine loads the
    page from a local origin with no network and no access to arbitrary paths —
    a linked image simply does not appear, which is how a CV would print with a
    blank square where the photo should be.
    """
    src = str(path or "")
    if not src or not os.path.isfile(src):
        return ""
    suffix = os.path.splitext(src)[1].lower()
    mime = _IMAGE_MIME.get(suffix, "")
    if not mime:
        return ""
    try:
        with open(src, "rb") as fh:
            raw = fh.read()
    except OSError as exc:
        obs.report("supercomputer.documents.image_read", exc, level="debug")
        return ""

    if len(raw) > MAX_IMAGE_BYTES or suffix == ".bmp":
        with obs.swallow("supercomputer.documents.image_resize", path=src):
            from PIL import Image

            buf = io.BytesIO()
            img = Image.open(io.BytesIO(raw))
            img.thumbnail((max_edge, max_edge))
            jpeg = suffix in (".jpg", ".jpeg")
            if img.mode in ("RGBA", "P") and jpeg:
                img = img.convert("RGB")
            img.save(buf, format="JPEG" if jpeg else "PNG")
            raw = buf.getvalue()
            mime = "image/jpeg" if jpeg else "image/png"
    if len(raw) > MAX_IMAGE_BYTES:
        return ""
    return f"data:{mime};base64,{base64.b64encode(raw).decode('ascii')}"


def embed_images(task_id: str, body_html: str) -> str:
    """Turn ``<img src="chart.png">`` into an embedded image.

    A chart the agent generated with run_python lands in the task workspace, and
    the model then writes an ordinary markdown image referencing it by name.
    This is what joins the two up — without it the PDF prints a broken-image box
    where the chart should be.
    """
    body = str(body_html or "")
    if "<img" not in body:
        return body

    def _swap(match: "re.Match") -> str:
        ref = match.group(2)
        if ref.startswith(("data:", "http://", "https://")):
            return match.group(0)
        target, _reason = workspace.resolve(task_id, ref)
        uri = image_data_uri(target) if target else ""
        return f'{match.group(1)}{uri or ref}{match.group(3)}' if uri else match.group(0)

    return re.sub(r'(<img\b[^>]*?\bsrc=")([^"]+)(")', _swap, body, flags=re.I)


# ── table parsing ──────────────────────────────────────────────────────────
def parse_rows(content: Any) -> List[List[str]]:
    """Coerce whatever the model produced into rows.

    Accepts a list of lists, a list of dicts (keys become the header row), or a
    markdown / CSV-ish table as text. Returning [] means "that was not a table",
    which the caller reports rather than writing an empty spreadsheet.
    """
    if isinstance(content, list) and content:
        if all(isinstance(r, dict) for r in content):
            headers: List[str] = []
            for row in content:
                for k in row:
                    if k not in headers:
                        headers.append(str(k))
            return [headers] + [[str(r.get(h, "")) for h in headers] for r in content]
        if all(isinstance(r, (list, tuple)) for r in content):
            return [[str(c) for c in r] for r in content]

    text = str(content or "").strip()
    if not text:
        return []
    # A markdown/CSV-ish table has a header AND at least one data row. Without
    # that floor, one line of prose split on its commas became a "spreadsheet"
    # of sentence fragments — a file that opens in Excel and says nothing.
    # An explicit list-of-lists or list-of-dicts (handled above) is always
    # honoured, because there the caller has said what it means.
    looks_tabular = sum(1 for ln in text.splitlines() if "|" in ln or "," in ln or "\t" in ln)
    if looks_tabular < 2:
        return []
    rows: List[List[str]] = []
    for line in text.splitlines():
        ln = line.strip()
        if not ln:
            continue
        # A markdown separator row (|---|---|) carries no data.
        if re.fullmatch(r"\|?[\s:|-]+\|?", ln) and "-" in ln:
            continue
        if "|" in ln:
            cells = [c.strip() for c in ln.strip("|").split("|")]
        else:
            try:
                cells = next(csv.reader([ln]))
            except Exception:
                cells = [ln]
        if any(c for c in cells):
            rows.append([str(c) for c in cells])
    return rows


# ── writers ────────────────────────────────────────────────────────────────
def _coerce_number(value: Any) -> Any:
    """Numbers written as text sort and total wrongly in Excel.

    A scraped table gives "1,234.5", "42%", "£19.99" — all strings. Left alone
    the spreadsheet is a picture of data rather than data, and the first thing
    anyone does with it (sort by that column, sum it) is wrong.
    """
    text = str(value or "").strip()
    if not text:
        return ""
    cleaned = text.replace(",", "").replace("£", "").replace("$", "").replace("€", "")
    percent = cleaned.endswith("%")
    if percent:
        cleaned = cleaned[:-1]
    try:
        number = float(cleaned)
    except ValueError:
        return text
    if percent:
        return number / 100.0
    return int(number) if number.is_integer() and "." not in cleaned else number


def _write_xlsx(task_id: str, rel: str, title: str, rows: List[List[str]],
                *, chart: bool = False) -> Dict[str, Any]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        obs.report("supercomputer.documents.openpyxl", exc)
        return {"ok": False, "reason": "xlsx_unavailable"}
    wb = Workbook()
    ws = wb.active
    ws.title = _slug(title)[:31] or "Sheet1"

    for index, row in enumerate(rows):
        ws.append(list(row) if index == 0 else [_coerce_number(c) for c in row])

    if rows:
        header_fill = PatternFill("solid", fgColor="1F2937")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 20
        # Freeze the header and turn on filtering: a spreadsheet you cannot
        # scroll without losing the column names is not much of a deliverable.
        ws.freeze_panes = "A2"
        if len(rows) > 1:
            ws.auto_filter.ref = ws.dimensions

        for i, _ in enumerate(rows[0], start=1):
            longest = max((len(str(r[i - 1])) for r in rows if len(r) >= i), default=10)
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
            # Percent-looking source values keep their meaning in the format.
            if any(str(r[i - 1]).strip().endswith("%") for r in rows[1:] if len(r) >= i):
                for cell in ws[get_column_letter(i)][1:]:
                    cell.number_format = "0.0%"

        if chart and len(rows) > 2 and len(rows[0]) > 1:
            with obs.swallow("supercomputer.documents.xlsx_chart"):
                from openpyxl.chart import BarChart, Reference

                bar = BarChart()
                bar.title = title
                bar.height, bar.width = 8, 18
                data = Reference(ws, min_col=2, min_row=1,
                                 max_col=len(rows[0]), max_row=len(rows))
                cats = Reference(ws, min_col=1, min_row=2, max_row=len(rows))
                bar.add_data(data, titles_from_data=True)
                bar.set_categories(cats)
                ws.add_chart(bar, f"{get_column_letter(len(rows[0]) + 2)}2")

    buf = io.BytesIO()
    wb.save(buf)
    return workspace.write_bytes(task_id, rel, buf.getvalue())


_MD_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _docx_runs(paragraph: Any, text: str) -> None:
    """Write text into a paragraph, honouring **bold**, *italic* and `code`.

    Without this every docx came out in one flat weight — the "no bold points"
    complaint, literally. python-docx has no markdown support, so the inline
    spans have to be split into runs by hand.
    """
    for part in _MD_INLINE.split(str(text or "")):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def _write_docx(task_id: str, rel: str, title: str, text: str,
                *, photo: str = "") -> Dict[str, Any]:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except Exception as exc:
        obs.report("supercomputer.documents.python_docx", exc)
        return {"ok": False, "reason": "docx_unavailable"}
    doc = Document()
    with obs.swallow("supercomputer.documents.docx_base_font"):
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

    doc.add_heading(str(title or "Document"), level=0)
    if photo and os.path.isfile(photo):
        with obs.swallow("supercomputer.documents.docx_photo", path=photo):
            doc.add_picture(photo, width=Inches(1.4))

    for block in re.split(r"\n{2,}", str(text or "")):
        chunk = block.strip()
        if not chunk:
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", chunk)
        if heading:
            doc.add_heading(heading.group(2), level=min(4, len(heading.group(1))))
            continue
        lines = chunk.splitlines()
        # A markdown table becomes a real Word table, not a row of pipes.
        if len(lines) >= 2 and all("|" in ln for ln in lines[:2]):
            rows = parse_rows(chunk)
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for index, row in enumerate(rows):
                    cells = table.add_row().cells
                    for col, value in enumerate(row[:len(rows[0])]):
                        cells[col].text = ""
                        run = cells[col].paragraphs[0].add_run(str(value))
                        if index == 0:
                            run.bold = True
                continue
        if all(re.match(r"^\s*[-*]\s+", ln) for ln in lines):
            for ln in lines:
                para = doc.add_paragraph(style="List Bullet")
                _docx_runs(para, re.sub(r"^\s*[-*]\s+", "", ln))
            continue
        if all(re.match(r"^\s*\d+[.)]\s+", ln) for ln in lines):
            for ln in lines:
                para = doc.add_paragraph(style="List Number")
                _docx_runs(para, re.sub(r"^\s*\d+[.)]\s+", "", ln))
            continue
        _docx_runs(doc.add_paragraph(), chunk)
    buf = io.BytesIO()
    doc.save(buf)
    return workspace.write_bytes(task_id, rel, buf.getvalue())


def _write_pptx(task_id: str, rel: str, title: str, text: str) -> Dict[str, Any]:
    """Slides from markdown: each `#`/`##` heading starts one, bullets fill it.

    python-pptx has been installed and unused since the DocIntel work. "Make me
    a deck" was answered with a markdown file.
    """
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except Exception as exc:
        obs.report("supercomputer.documents.python_pptx", exc)
        return {"ok": False, "reason": "pptx_unavailable"}

    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = str(title or "Presentation")
    with obs.swallow("supercomputer.documents.pptx_subtitle"):
        cover.placeholders[1].text = time.strftime("%d %B %Y")

    slide = None
    body = None
    for raw in str(text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,3}\s+(.*)$", line)
        if heading:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = heading.group(1)[:120]
            body = slide.placeholders[1].text_frame
            body.clear()
            continue
        if slide is None:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = str(title or "Notes")[:120]
            body = slide.placeholders[1].text_frame
            body.clear()
        para = body.add_paragraph() if body.text else body.paragraphs[0]
        para.text = re.sub(r"^\s*[-*\d.)]+\s*", "", line)[:220]
        para.font.size = Pt(18)

    buf = io.BytesIO()
    deck.save(buf)
    return workspace.write_bytes(task_id, rel, buf.getvalue())


def _write_pdf(task_id: str, rel: str, title: str, body_html: str,
               *, config: Optional[Dict[str, Any]] = None,
               template: str = "", theme: str = "", subtitle: str = "",
               photo: str = "", meta: Optional[Dict[str, Any]] = None,
               landscape: bool = False) -> Dict[str, Any]:
    """HTML → local preview server → browser tab → printToPDF → workspace.

    The intermediate .html is kept: it is the source of the PDF, and when the
    print step fails it is a real deliverable on its own rather than nothing.
    """
    from . import preview
    from .browser import controller as browser

    html_rel = re.sub(r"\.pdf$", "", rel, flags=re.I) + ".html"
    document = wrap_html(title, embed_images(task_id, body_html),
                         template=template, theme=theme, subtitle=subtitle,
                         photo=photo, meta=meta, config=config)
    wrote = workspace.write(task_id, html_rel, document)
    if not wrote.get("ok"):
        return wrote

    started = preview.start(task_id)
    if not started.get("ok"):
        return {"ok": False, "reason": f"preview_unavailable:{started.get('reason')}",
                "data": {"html": (wrote.get("data") or {}).get("path", "")}}
    base = str((started.get("data") or {}).get("url") or "").rstrip("/")

    opened = browser.open_url(f"{base}/{html_rel}", new_tab=True, config=config)
    if not opened.get("ok"):
        return {"ok": False, "reason": f"preview_open_failed:{opened.get('reason')}",
                "data": {"html": (wrote.get("data") or {}).get("path", "")}}

    # Give the page a moment to lay out — a template with an embedded photo or
    # a chart prints blank if the print fires before the image decodes.
    browser.wait(0.7)
    printed = browser.print_pdf(config=config, landscape=bool(landscape))
    if not printed.get("ok"):
        return {"ok": False, "reason": f"print_failed:{printed.get('reason')}",
                "data": {"html": (wrote.get("data") or {}).get("path", "")}}
    b64 = str((printed.get("data") or {}).get("pdf_base64") or "")
    if not b64:
        return {"ok": False, "reason": "print_returned_nothing"}
    try:
        blob = base64.b64decode(b64)
    except Exception as exc:
        obs.report("supercomputer.documents.pdf_decode", exc)
        return {"ok": False, "reason": "pdf_decode_failed"}
    return workspace.write_bytes(task_id, rel, blob)


_FORMAT_ALIASES = {
    "markdown": "md", "text": "txt", "excel": "xlsx", "spreadsheet": "xlsx",
    "xls": "xlsx", "word": "docx", "doc": "docx", "slides": "pptx",
    "powerpoint": "pptx", "deck": "pptx", "presentation": "pptx",
    "webpage": "html", "web": "html",
}


def normalise_format(fmt: str) -> str:
    f = str(fmt or "md").strip().lower().lstrip(".")
    return _FORMAT_ALIASES.get(f, f)


def create(
    task_id: str,
    name: str,
    fmt: str,
    content: Any,
    *,
    title: str = "",
    subtitle: str = "",
    template: str = "",
    theme: str = "",
    photo: str = "",
    meta: Optional[Dict[str, Any]] = None,
    landscape: bool = False,
    config: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Write one document. Returns the standard ``{ok, reason, data}`` envelope.

    ``template`` and ``theme`` apply to the laid-out formats (html and pdf). If
    neither is named they are inferred from the document's name and title, so
    "my CV" gets the CV layout rather than the generic report shell — the model
    does not have to know the template system exists to benefit from it.
    """
    f = normalise_format(fmt)
    if f not in FORMATS:
        return {"ok": False, "reason": f"unsupported_format:{f}"}

    rel = filename_for(name, f)
    heading = str(title or name or "Document")
    picked = doc_templates.resolve(template, hint=f"{name} {title} {subtitle}")
    picked_theme = doc_themes.resolve(theme, template=picked, config=config)
    photo_path = str(photo or "")
    extras = {"template": picked, "theme": picked_theme}

    if f in ("md", "txt"):
        out = workspace.write(task_id, rel, str(content or ""))
    elif f == "html":
        body = str(content or "")
        # A full document from the model is used as-is; a fragment or markdown
        # goes through the template so it prints and reads properly.
        if "<html" in body.lower() or "<!doctype" in body.lower():
            out = workspace.write(task_id, rel, body)
        else:
            out = workspace.write(task_id, rel, wrap_html(
                heading, embed_images(task_id, markdown_to_html(body)),
                template=picked, theme=picked_theme, subtitle=subtitle,
                photo=image_data_uri(photo_path) if photo_path else "",
                meta=meta, config=config,
            ))
    elif f == "csv":
        rows = parse_rows(content)
        if not rows:
            return {"ok": False, "reason": "no_tabular_data"}
        buf = io.StringIO(newline="")
        csv.writer(buf, lineterminator="\n").writerows(rows)
        out = workspace.write(task_id, rel, buf.getvalue())
    elif f == "xlsx":
        rows = parse_rows(content)
        if not rows:
            return {"ok": False, "reason": "no_tabular_data"}
        out = _write_xlsx(task_id, rel, heading, rows,
                          chart=bool((meta or {}).get("chart")))
    elif f == "docx":
        out = _write_docx(task_id, rel, heading, str(content or ""), photo=photo_path)
    elif f == "pptx":
        out = _write_pptx(task_id, rel, heading, str(content or ""))
    else:
        body = str(content or "")
        body_html = (body if ("<p>" in body or "<h1" in body or "<table" in body)
                     else markdown_to_html(body))
        out = _write_pdf(task_id, rel, heading, body_html, config=config,
                         template=picked, theme=picked_theme, subtitle=subtitle,
                         photo=image_data_uri(photo_path) if photo_path else "",
                         meta=meta, landscape=landscape)

    # Report the choices back, so the agent can say which template and theme it
    # picked rather than leave the result unexplained.
    if isinstance(out, dict) and out.get("ok"):
        out.setdefault("data", {})
        if isinstance(out["data"], dict) and f in ("pdf", "html"):
            out["data"].update(extras)
    return out


def describe(fmt: str) -> str:
    return {
        "md": "a markdown document", "txt": "a text file",
        "html": "a web page", "csv": "a CSV spreadsheet",
        "xlsx": "an Excel spreadsheet", "docx": "a Word document",
        "pptx": "a slide deck", "pdf": "a PDF",
    }.get(normalise_format(fmt), "a document")


def describe_options() -> str:
    """Templates and themes, for the tool description the model reads."""
    return (f"TEMPLATES: {doc_templates.describe()}.\n"
            f"THEMES: {doc_themes.describe()}.")


__all__ = ["FORMATS", "create", "describe", "describe_options", "filename_for",
           "parse_rows", "markdown_to_html", "wrap_html", "normalise_format",
           "image_data_uri", "embed_images"]
