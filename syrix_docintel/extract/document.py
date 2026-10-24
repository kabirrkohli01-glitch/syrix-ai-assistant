"""Word / rich-text document extraction — DOCX, RTF, ODT, (legacy DOC).

DOCX is the primary path (python-docx): headings drive a section hierarchy,
body text is grouped under its heading, and tables become TableObjects. RTF and
ODT are best-effort text extraction (fast-follow); legacy binary .doc is
refused with a clear note rather than mangled.
"""
from __future__ import annotations

import io
from typing import List

import syrix_obs

from .. import types as T
from ..chunk import add_page, finalize
from ..safety import Deadline, ProcessingError
from ..types import DocumentObject, TableObject


def extract(blob: bytes, doc: DocumentObject, deadline: Deadline) -> None:
    if doc.subtype == "docx":
        _extract_docx(blob, doc)
    elif doc.subtype == "rtf":
        _extract_rtf(blob, doc)
    elif doc.subtype == "odt":
        _extract_odt(blob, doc)
    elif doc.subtype in ("doc", "ole"):
        raise ProcessingError("Legacy binary .doc isn't supported yet — please save as .docx or PDF.")
    else:
        _extract_rtf(blob, doc)  # last-ditch: try as rich text
    finalize(doc)


def _heading_level(style_name: str) -> int:
    s = (style_name or "").lower()
    if s.startswith("heading"):
        try:
            return int(s.replace("heading", "").strip() or "1")
        except Exception:
            return 1
    if s in ("title",):
        return 0
    return -1


def _extract_docx(blob: bytes, doc: DocumentObject) -> None:
    try:
        import docx  # type: ignore
    except Exception as exc:
        raise ProcessingError(f"DOCX support unavailable (install python-docx) — {exc}.")
    try:
        d = docx.Document(io.BytesIO(blob))
    except Exception as exc:
        raise ProcessingError(f"Could not open DOCX ({exc}).")
    doc.extraction_method = "python-docx"

    # Metadata from core properties.
    with syrix_obs.swallow("docintel.docx.core_props"):
        cp = d.core_properties
        for attr in ("author", "title", "created", "modified", "last_modified_by", "subject"):
            v = getattr(cp, attr, None)
            if v:
                doc.metadata[attr] = str(v)[:200]

    # Section-grouped paragraphs.
    section_path: List[str] = []
    buf: List[str] = []
    cur_section = ""

    def flush():
        nonlocal buf
        if buf:
            add_page(doc, 0, cur_section, "\n".join(buf), kind="text")
            buf = []

    for para in d.paragraphs:
        text = (para.text or "").strip()
        lvl = _heading_level(getattr(getattr(para, "style", None), "name", "") or "")
        if lvl >= 0 and text:
            flush()
            # Maintain a heading path stack keyed by level.
            section_path = section_path[:lvl]
            section_path.append(text)
            cur_section = " › ".join(section_path)
            buf.append(text)   # keep the heading text in-context
        elif text:
            buf.append(text)
    flush()

    # Tables.
    for ti, table in enumerate(d.tables):
        try:
            rows = [[(c.text or "").strip() for c in row.cells] for row in table.rows]
        except Exception:
            continue
        if not rows:
            continue
        cols = rows[0] if rows else []
        body = rows[1:] if len(rows) > 1 else []
        tobj = TableObject(page=0, title=f"Table {ti + 1}", columns=cols,
                           rows=body[:200], n_rows=len(body), n_cols=len(cols))
        doc.tables.append(tobj)
        add_page(doc, 0, f"Table {ti + 1}", tobj.preview_md(max_rows=15), kind="table")

    doc.confidence = 1.0
    if not doc.full_text().strip():
        doc.status = T.STATUS_PARTIAL
        doc.add_note("No readable text found in the document.")


def _extract_rtf(blob: bytes, doc: DocumentObject) -> None:
    try:
        from striprtf.striprtf import rtf_to_text  # type: ignore
    except Exception as exc:
        raise ProcessingError(f"RTF support unavailable (install striprtf) — {exc}.")
    try:
        text = rtf_to_text(blob.decode("utf-8", errors="ignore"))
    except Exception as exc:
        raise ProcessingError(f"Could not parse RTF ({exc}).")
    doc.extraction_method = "striprtf"
    add_page(doc, 0, "", text, kind="text")
    doc.confidence = 1.0


def _extract_odt(blob: bytes, doc: DocumentObject) -> None:
    try:
        from odf.opendocument import load  # type: ignore
        from odf import teletype  # type: ignore
        from odf.text import P  # type: ignore
    except Exception as exc:
        raise ProcessingError(f"ODT support unavailable (install odfpy) — {exc}.")
    try:
        odt = load(io.BytesIO(blob))
        paras = odt.getElementsByType(P)
        text = "\n".join(teletype.extractText(p) for p in paras)
    except Exception as exc:
        raise ProcessingError(f"Could not parse ODT ({exc}).")
    doc.extraction_method = "odfpy"
    add_page(doc, 0, "", text, kind="text")
    doc.confidence = 1.0
