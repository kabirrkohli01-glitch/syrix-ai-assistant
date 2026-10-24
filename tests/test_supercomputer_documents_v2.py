"""Templates, themes and images — the "these PDFs are lazy" fix.

The old output went through exactly one hardcoded stylesheet, so a CV came out
looking like a memo and nothing looked designed. These assert the things that
made it lazy are now impossible: one shell for everything, no emphasis, no
images, no page numbers.
"""
from __future__ import annotations

import os
import sys
import unittest

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
for _p in (_ROOT, _HERE):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from supercomputer_testkit import SupercomputerTestMixin  # noqa: E402

from syrix_supercomputer import doc_templates, doc_themes, documents  # noqa: E402

_MD = """# Experience
**Senior Engineer** at Acme
- Built *things*
- Shipped `code`

| Item | Value |
|---|---|
| Speed | 42% |
| Cost | £1,299 |
"""


class TestTemplatesAndThemes(SupercomputerTestMixin, unittest.TestCase):
    def test_every_template_renders_with_every_theme(self):
        """Template x theme is a real matrix, not a promise."""
        body = documents.markdown_to_html(_MD)
        for template in doc_templates.names():
            for theme in doc_themes.names():
                html = doc_templates.render(body, template=template, theme=theme,
                                            title="Kabirr Kohli", subtitle="Engineer")
                self.assertIn("<!doctype html>", html.lower(), f"{template}/{theme}")
                self.assertIn("Kabirr Kohli", html, f"{template}/{theme}")
                self.assertIn(doc_themes.THEMES[theme]["accent"], html,
                              f"{template}/{theme} did not apply the theme")

    def test_documents_are_self_contained(self):
        """The print runs against a local origin with NO network, so anything
        remote silently fails to render."""
        html = doc_templates.render("<p>x</p>", template="report", title="T")
        for remote in ("http://", "https://", "//fonts.", "@import"):
            self.assertNotIn(remote, html, f"a remote reference would not render: {remote}")

    def test_pages_are_numbered(self):
        html = doc_templates.render("<p>x</p>", template="report", title="T")
        self.assertIn("counter(page)", html)

    def test_headings_stay_with_their_text(self):
        html = doc_templates.render("<p>x</p>", template="report", title="T")
        self.assertIn("break-after: avoid", html)

    def test_the_template_is_inferred_when_nobody_names_one(self):
        self.assertEqual(doc_templates.resolve("", hint="my CV"), "cv")
        self.assertEqual(doc_templates.resolve("", hint="pros and cons of X"), "comparison")
        self.assertEqual(doc_templates.resolve("", hint="Invoice 0042"), "invoice")
        self.assertEqual(doc_templates.resolve("", hint="NVDA trade setup"), "trading_setup")
        self.assertEqual(doc_templates.resolve("", hint="covering letter"), "letter")
        self.assertEqual(doc_templates.resolve("", hint="something vague"), "report")

    def test_an_unknown_template_or_theme_falls_back_rather_than_failing(self):
        html = doc_templates.render("<p>x</p>", template="nonsense", theme="chartreuse",
                                    title="T")
        self.assertIn("<!doctype html>", html.lower())

    def test_the_model_is_told_what_it_can_choose_from(self):
        described = documents.describe_options()
        for template in doc_templates.names():
            self.assertIn(template, described)
        for theme in doc_themes.names():
            self.assertIn(theme, described)


class TestFormats(SupercomputerTestMixin, unittest.TestCase):
    def test_every_declared_format_actually_writes_a_file(self):
        for fmt in ("md", "txt", "html", "csv", "xlsx", "docx", "pptx"):
            res = documents.create("t1", f"sample_{fmt}", fmt, _MD,
                                   title="Kabirr Kohli", subtitle="Engineer")
            self.assertTrue(res["ok"], f"{fmt}: {res.get('reason')}")
            path = res["data"]["path"]
            self.assertTrue(os.path.isfile(path), fmt)
            self.assertGreater(os.path.getsize(path), 40, fmt)

    def test_format_aliases_are_accepted(self):
        for alias, real in (("excel", "xlsx"), ("word", "docx"), ("slides", "pptx"),
                            ("markdown", "md"), ("presentation", "pptx")):
            self.assertEqual(documents.normalise_format(alias), real)

    def test_a_word_document_carries_real_emphasis_and_tables(self):
        """The "no bold points" complaint, literally. python-docx has no
        markdown support, so inline spans have to be split into runs by hand."""
        res = documents.create("t1", "cv", "docx", _MD, title="Kabirr")
        self.assertTrue(res["ok"], res.get("reason"))
        from docx import Document

        doc = Document(res["data"]["path"])
        self.assertTrue(doc.tables, "a markdown table must become a real table")
        bolded = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
        self.assertIn("Senior Engineer", bolded)
        italics = [r.text for p in doc.paragraphs for r in p.runs if r.italic]
        self.assertIn("things", italics)

    def test_html_reports_which_template_and_theme_it_used(self):
        """So the agent can say which template and theme it picked."""
        res = documents.create("t1", "my cv", "html", _MD, title="Kabirr")
        self.assertTrue(res["ok"])
        self.assertEqual(res["data"]["template"], "cv")
        self.assertIn(res["data"]["theme"], doc_themes.names())

    def test_an_unsupported_format_is_refused_by_name(self):
        res = documents.create("t1", "x", "epub", "body")
        self.assertFalse(res["ok"])
        self.assertEqual(res["reason"], "unsupported_format:epub")

    def test_a_spreadsheet_with_no_table_is_refused_not_written_empty(self):
        res = documents.create("t1", "x", "xlsx", "just some prose, no table here")
        self.assertFalse(res["ok"])
        self.assertEqual(res["reason"], "no_tabular_data")


class TestImages(SupercomputerTestMixin, unittest.TestCase):
    def _png(self, name="photo.png"):
        from PIL import Image

        path = os.path.join(self._tmp, name)
        Image.new("RGB", (64, 64), (200, 30, 30)).save(path)
        return path

    def test_an_image_becomes_a_data_uri(self):
        """A linked file simply does not render in the print origin — that is
        how a CV would print with a blank square where the photo should be."""
        uri = documents.image_data_uri(self._png())
        self.assertTrue(uri.startswith("data:image/png;base64,"))

    def test_a_missing_or_non_image_file_yields_nothing_rather_than_breaking(self):
        self.assertEqual(documents.image_data_uri(""), "")
        self.assertEqual(documents.image_data_uri(os.path.join(self._tmp, "nope.png")), "")
        text = os.path.join(self._tmp, "notes.txt")
        with open(text, "w", encoding="utf-8") as fh:
            fh.write("hello")
        self.assertEqual(documents.image_data_uri(text), "")

    def test_a_photo_reaches_the_cv_header(self):
        html = doc_templates.render("<p>x</p>", template="cv", title="Kabirr",
                                    photo=documents.image_data_uri(self._png()))
        self.assertIn("sc-photo", html)
        self.assertIn("data:image/png;base64,", html)

    def test_a_workspace_chart_is_embedded_into_the_document(self):
        from syrix_supercomputer import workspace

        from PIL import Image

        target, _ = workspace.resolve("t1", "chart.png")
        os.makedirs(os.path.dirname(target), exist_ok=True)
        Image.new("RGB", (40, 40), (10, 90, 200)).save(target)
        out = documents.embed_images("t1", '<img src="chart.png" alt="">')
        self.assertIn("data:image/png;base64,", out)

    def test_a_remote_image_reference_is_left_alone(self):
        body = '<img src="https://example.com/x.png">'
        self.assertEqual(documents.embed_images("t1", body), body)


if __name__ == "__main__":
    unittest.main()
