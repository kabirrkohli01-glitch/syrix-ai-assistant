"""Document Intelligence upgrade pass (2026-07-16): follow-up session context,
edited-version export, spreadsheet data-quality findings, grounding document
map + OCR uncertainty, and the PaddleOCR 3.x result parser."""
import io
import os
import sys
import time
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest


@pytest.fixture(autouse=True)
def _env():
    import core_boot
    core_boot.CONFIG["docintel_persist"] = False
    core_boot.CONFIG["docintel_use_embeddings"] = False
    core_boot.CONFIG["docintel_ocr_enabled"] = False
    core_boot.CONFIG["docintel_followup_enabled"] = True
    import syrix_docintel.registry as reg
    reg._reset_for_tests()
    import syrix_docintel.session as sess
    sess._reset_for_tests()
    yield
    sess._reset_for_tests()


def _analyze(name, blob):
    import syrix_docintel as di
    return di.analyze_upload(name, blob)


def _budget_xlsx(with_error=True, second_sheet=True):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["month", "spend"])
    for i, v in enumerate([100, 110, 120, 140, 150, 160, 170, 180, 190, 200], 1):
        ws.append([f"M{i}", v])
    ws.append(["M1", 100])   # duplicate of the first data row
    if with_error:
        ws["B3"] = "#DIV/0!"
    if second_sheet:
        ws2 = wb.create_sheet("Summary")
        ws2["A1"] = "=Budget!B2+Budget!B3"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# ── spreadsheet quality ─────────────────────────────────────────────────────

def test_spreadsheet_quality_findings():
    d = _analyze("budget.xlsx", _budget_xlsx())
    q = d.statistics["quality"]["Budget"]
    assert any("#DIV/0!" in e for e in q["formula_errors"])
    assert q["duplicate_rows"] == 1
    assert q["missing_cells"]["spend"] == 1     # the error cell reads as NaN
    # The findings must be retrievable text, not just metadata.
    assert any("Data quality findings" in c.text for c in d.chunks)


def test_spreadsheet_trend_and_sheet_refs():
    d = _analyze("budget.xlsx", _budget_xlsx(with_error=False))
    assert d.statistics["sheets"]["Budget"]["spend"]["trend"] == "rising"
    assert d.metadata["sheet_refs"] == {"Summary": ["Budget"]}


def test_spreadsheet_macro_flag_never_executes():
    blob = _budget_xlsx(with_error=False, second_sheet=False)
    src = io.BytesIO(blob)
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/vbaProject.bin", b"\x00fake-macro-blob")
    d = _analyze("macro.xlsm", out.getvalue())
    assert d.metadata.get("has_macros") is True
    assert any("NOT executed" in n for n in d.notes)
    assert d.status in ("ready", "partial")


def test_tsv_reads_as_spreadsheet():
    rows = "n\tval\n" + "\n".join(f"{i}\t{i * 10}" for i in range(1, 6))
    d = _analyze("data.tsv", rows.encode())
    assert d.kind in ("spreadsheet", "text")
    if d.kind == "spreadsheet":
        assert d.statistics["sheets"]["Sheet1"]["val"]["max"] == 50.0


# ── grounding: document map + formulas + OCR uncertainty ───────────────────

def test_grounding_has_document_map_and_formulas():
    import syrix_docintel as di
    d = _analyze("budget.xlsx", _budget_xlsx())
    ctx, cites = di.build_grounding([d.id], "Explain the formula in A1 of the summary")
    assert "--- FILE: budget.xlsx ---" in ctx
    assert "Sheets:" in ctx
    assert "=Budget!B2+Budget!B3" in ctx          # formula surfaced for a cell ask
    assert any("budget.xlsx" in c for c in cites)


def test_grounding_flags_ocr_uncertainty():
    import syrix_docintel as di
    from syrix_docintel import registry
    from syrix_docintel.types import ChunkObject, DocumentObject
    doc = DocumentObject(id="doc-ocr", name="scan.pdf", kind="pdf", pages=1)
    doc.metadata["ocr_used"] = True
    doc.metadata["ocr_confidence"] = 0.62
    doc.chunks = [ChunkObject(id="doc-ocr::c0", text="Total due 4200 USD", page=1, kind="ocr")]
    registry.put(doc)
    ctx, _ = di.build_grounding(["doc-ocr"], "what is the total due?")
    assert "OCR" in ctx and "62%" in ctx
    assert "may contain" in ctx.lower()


def test_grounding_multi_file_compare_rule():
    import syrix_docintel as di
    a = _analyze("a.txt", b"Alpha report: revenue was 10 million dollars in 2025.")
    b = _analyze("b.txt", b"Beta report: revenue was 12 million dollars in 2025.")
    ctx, cites = di.build_grounding([a.id, b.id], "compare the revenue in the two files")
    assert "attribute every claim" in ctx
    assert any("a.txt" in c for c in cites) and any("b.txt" in c for c in cites)


# ── follow-up session context ───────────────────────────────────────────────

def test_followup_resolves_explicit_reference():
    from syrix_docintel import session
    d = _analyze("report.txt", b"Quarterly revenue rose 18 percent to 4.2M.")
    session.note_used([d.id])
    assert session.resolve_followup("what does the document say about revenue?") == [d.id]
    assert session.resolve_followup("summarize the pdf again") == [d.id]
    assert session.resolve_followup("is there anything about margins in the file?") == [d.id]


def test_followup_resolves_filename_mention():
    from syrix_docintel import session
    d = _analyze("q3-financials.txt", b"Numbers for the third quarter.")
    session.note_used([d.id])
    assert session.resolve_followup("what's in q3-financials?") == [d.id]


def test_followup_continuation_window():
    from syrix_docintel import session
    d = _analyze("report.txt", b"Quarterly revenue rose 18 percent.")
    session.note_used([d.id])
    # Content follow-up without a doc word — allowed inside the window.
    assert session.resolve_followup("explain that in simpler words") == [d.id]
    # Clearly unrelated ask — never grounded.
    assert session.resolve_followup("what's the weather in Delhi?") == []
    assert session.resolve_followup("open chrome") == []


def test_followup_ttl_and_kill_switch():
    import core_boot
    from syrix_docintel import session
    d = _analyze("report.txt", b"Quarterly revenue rose 18 percent.")
    session.note_used([d.id])
    session._STATE["ts"] = time.time() - 99999
    assert session.resolve_followup("what does the document say?") == []
    session.note_used([d.id])
    core_boot.CONFIG["docintel_followup_enabled"] = False
    try:
        assert session.resolve_followup("what does the document say?") == []
    finally:
        core_boot.CONFIG["docintel_followup_enabled"] = True


def test_followup_empty_without_docs():
    from syrix_docintel import session
    assert session.resolve_followup("what does the document say?") == []


# ── edited-version export ───────────────────────────────────────────────────

def test_export_intent_detection():
    from syrix_docintel import export
    assert export.detect_export_intent("rewrite this document and give me a downloadable version")
    assert export.detect_export_intent("improve the language and create an edited copy")
    assert export.detect_export_intent("fix the grammar and save a new version")
    assert not export.detect_export_intent("summarize this document")
    assert not export.detect_export_intent("what are the key points?")
    assert not export.detect_export_intent("is there anything about revenue?")


def test_export_writes_docx_for_document(tmp_path, monkeypatch):
    from syrix_docintel import export
    monkeypatch.setattr(export, "export_dir", lambda: str(tmp_path))
    d = _analyze("notes.docx.txt", b"placeholder")   # just need a DocumentObject
    d.kind = "document"
    d.name = "notes.docx"
    path, err = export.write_export(d, "# Title\n\n- point one\n\nBody paragraph.")
    assert err == "" and path and path.endswith(".docx")
    assert os.path.exists(path)
    import docx
    dd = docx.Document(path)
    texts = [p.text for p in dd.paragraphs]
    assert "Title" in texts and "point one" in texts and "Body paragraph." in texts


def test_export_writes_text_for_text_source(tmp_path, monkeypatch):
    from syrix_docintel import export
    monkeypatch.setattr(export, "export_dir", lambda: str(tmp_path))
    d = _analyze("plain.txt", b"original")
    path, err = export.write_export(d, "improved text")
    assert err == "" and path and path.endswith(".txt")
    with open(path, encoding="utf-8") as f:
        assert f.read().strip() == "improved text"


def test_export_rejects_empty_and_sanitises_name(tmp_path, monkeypatch):
    from syrix_docintel import export
    monkeypatch.setattr(export, "export_dir", lambda: str(tmp_path))
    d = _analyze("evil.txt", b"x")
    d.name = "..\\..\\evil<>:name.txt"
    path, err = export.write_export(d, "")
    assert path is None and err
    path, err = export.write_export(d, "content")
    assert path and os.path.dirname(path) == str(tmp_path)
    assert ".." not in os.path.basename(path)


def test_export_prompt_keeps_structure_rules():
    from syrix_docintel import export
    d = _analyze("doc.txt", b"Alpha.\n\nBeta.")
    p = export.build_rewrite_prompt(d, "keep the same document and only improve its language")
    assert "Change ONLY the language" in p
    assert "ONLY the rewritten document body" in p
    assert "Alpha." in p


# ── OCR result parsing (both PaddleOCR API generations) ─────────────────────

def test_paddle_result_parser_v3_and_v2():
    from syrix_docintel.ocr import _parse_paddle_result
    v3 = [{"rec_texts": ["Hello", "World"], "rec_scores": [0.98, 0.91]}]
    lines, scores = _parse_paddle_result(v3)
    assert lines == ["Hello", "World"] and scores == [0.98, 0.91]
    v2 = [[[[0, 0], ("Legacy", 0.77)]]]
    lines2, scores2 = _parse_paddle_result(v2)
    assert lines2 == ["Legacy"] and scores2 == [0.77]
    assert _parse_paddle_result(None) == ([], [])


# ── presentation now actually extracts (python-pptx installed) ──────────────

def test_pptx_extracts_slides():
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Growth Strategy"
    s.placeholders[1].text = "Expand into two new markets"
    b = io.BytesIO()
    prs.save(b)
    d = _analyze("deck.pptx", b.getvalue())
    assert d.kind == "presentation" and d.status == "ready" and d.slides == 1
    assert "Growth Strategy" in d.full_text()
    assert d.chunks[0].page == 1     # citations resolve to Slide 1
